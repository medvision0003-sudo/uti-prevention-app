from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.dependencies import get_current_user
from app.models import BladderLog, ClinicIntake, DailyCheckIn, PatientIdentity, User, WaterLog


router = APIRouter(prefix="/api/export", tags=["Patient Data Export"])
INDIA_TIMEZONE = ZoneInfo("Asia/Kolkata")


def excel_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def local_datetime(value: datetime) -> datetime:
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(INDIA_TIMEZONE)


def prepare_sheet(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    sheet.freeze_panes = "A2"


def format_sheet(sheet) -> None:
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275B49")
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 40)
        sheet.column_dimensions[column[0].column_letter].width = width


@router.get("/excel")
def export_patient_excel(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    identity = db.scalar(select(PatientIdentity).where(PatientIdentity.user_id == user.id))
    if identity is None:
        identity = PatientIdentity(
            user_id=user.id,
            patient_id=f"PT-{user.id:010d}",
            full_name="",
            contact_number="",
        )
        db.add(identity)
        db.commit()
        db.refresh(identity)

    water_entries = db.scalars(
        select(WaterLog).where(WaterLog.user_id == user.id).order_by(WaterLog.logged_at)
    ).all()
    bladder_entries = db.scalars(
        select(BladderLog).where(BladderLog.user_id == user.id).order_by(BladderLog.logged_at)
    ).all()
    check_ins = db.scalars(
        select(DailyCheckIn)
        .where(DailyCheckIn.user_id == user.id)
        .order_by(DailyCheckIn.check_in_date)
    ).all()
    intake = db.scalar(select(ClinicIntake).where(ClinicIntake.user_id == user.id))

    workbook = Workbook()
    patient_sheet = workbook.active
    patient_sheet.title = "Patient"
    prepare_sheet(patient_sheet, ["Field", "Value"])
    for label, value in [
        ("Patient ID", identity.patient_id),
        ("Name", identity.full_name),
        ("Contact number", identity.contact_number),
        ("Email", user.email),
        ("Account created (UTC)", user.created_at.isoformat()),
    ]:
        patient_sheet.append([label, excel_text(value)])
    format_sheet(patient_sheet)

    intake_sheet = workbook.create_sheet("Clinic Intake")
    prepare_sheet(intake_sheet, ["Field", "Value"])
    if intake is not None:
        for label, value in [
            ("Age", intake.age),
            ("Last menstrual period", intake.last_menstrual_period.isoformat() if intake.last_menstrual_period else ""),
            ("Postmenopausal", intake.postmenopausal),
            ("History of hysterectomy", intake.hysterectomy),
            ("Approximate menopause duration", intake.menopause_duration or ""),
            ("Recurrent UTI", intake.recurrent_uti),
            ("Diabetes", intake.diabetes),
            ("Constipation", intake.constipation),
            ("History of loose motions", intake.loose_motions),
            ("Recent travel", intake.recent_travel),
            ("Other urological history", intake.other_urological_history),
            ("Prior intervention or procedure", intake.prior_intervention),
            ("Prior stone surgery", intake.prior_stone_surgery),
            ("Note", excel_text(intake.note or "")),
            ("Last updated (UTC)", intake.updated_at.isoformat()),
        ]:
            intake_sheet.append([label, value])
    format_sheet(intake_sheet)

    water_sheet = workbook.create_sheet("Water Log")
    prepare_sheet(water_sheet, ["Date", "Time (India)", "Amount (ml)"])
    for entry in water_entries:
        logged_at = local_datetime(entry.logged_at)
        water_sheet.append([logged_at.date().isoformat(), logged_at.strftime("%H:%M:%S"), entry.amount_ml])
    format_sheet(water_sheet)

    bladder_sheet = workbook.create_sheet("Bladder Log")
    prepare_sheet(bladder_sheet, ["Date", "Time (India)"])
    for entry in bladder_entries:
        logged_at = local_datetime(entry.logged_at)
        bladder_sheet.append([logged_at.date().isoformat(), logged_at.strftime("%H:%M:%S")])
    format_sheet(bladder_sheet)

    check_in_sheet = workbook.create_sheet("Daily Check-ins")
    prepare_sheet(check_in_sheet, ["Date", "Bowel status", "Symptoms", "Last updated (UTC)"])
    for entry in check_ins:
        check_in_sheet.append([
            entry.check_in_date.isoformat(),
            entry.bowel_status,
            ", ".join(entry.symptoms or []),
            entry.updated_at.isoformat(),
        ])
    format_sheet(check_in_sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bladderwise-{identity.patient_id}.xlsx"'},
    )
